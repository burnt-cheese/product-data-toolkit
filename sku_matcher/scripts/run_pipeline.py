# -*- coding: utf-8 -*-
"""
sku_matcher（单步版）
====================
    python run_pipeline.py match [货号表.xlsx]
    或: python run_pipeline.py [货号表.xlsx]   （默认即 match）
    → 货号匹配 SAP，输出 output\sku_match_result.xlsx
"""
import os, sys, subprocess
from openpyxl import load_workbook

BASE = os.path.dirname(os.path.abspath(__file__))
PROJ = os.path.dirname(BASE)
PY   = sys.executable
GEN  = os.path.join(BASE, "generate_matched.py")
OUT  = os.path.join(PROJ, "output")
RES  = os.path.join(OUT, "sku_match_result.xlsx")

def run(script, *args):
    """流式执行子脚本：边跑边打印，让用户看到 SAP 解析等耗时步骤的进度。
    解决「脚本无输出以为卡住」的体验问题。"""
    env = os.environ.copy()
    # 强制 Python 子进程无缓冲，子脚本里的 print(flush=True) 也能立刻到达终端
    env["PYTHONUNBUFFERED"] = "1"
    proc = subprocess.Popen(
        [PY, script, *args],
        stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
        encoding="utf-8", errors="replace",
        env=env,
    )
    for line in proc.stdout:
        # 实时打印到主进程 stdout
        sys.stdout.write(line)
        sys.stdout.flush()
    proc.wait()
    return proc.returncode == 0

def count_unmatched():
    try:
        wb = load_workbook(RES, data_only=True)
        if "未匹配货号" not in wb.sheetnames:
            return 0
        ws = wb["未匹配货号"]
        return max(ws.max_row - 1, 0)
    except Exception:
        return 0

def cmd_match(inp):
    print("=" * 62)
    print("  货号匹配（SAP）")
    print("=" * 62)
    print(f"输入货号表: {inp}")
    if not run(GEN, inp):
        print("[错误] 匹配失败"); sys.exit(1)
    n_unc = count_unmatched()
    if n_unc > 0:
        print(f"\n>>> 有 {n_unc} 个货号未匹配（SAP 中无此货号）")
        print("    请直接在 output\\sku_match_result.xlsx 的「未匹配货号」sheet 对应行补充：")
        print("    品名/条形码/采购单价/供应商 等信息，保存后再跑一次本工具即可。")
    else:
        print("\n>>> 全部货号均已匹配，无需人工补充")

def main():
    args = sys.argv[1:]
    # 第一步 match：python run_pipeline.py match [货号表.xlsx]
    # 简化调用：python run_pipeline.py [货号表.xlsx]
    inp = os.path.join(PROJ, "input", "货号信息.xlsx")
    if args:
        if args[0] == "match":
            if len(args) > 1:
                inp = args[1]
        else:
            inp = args[0]
    cmd_match(inp)

if __name__ == "__main__":
    main()
