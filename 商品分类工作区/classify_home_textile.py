# -*- coding: utf-8 -*-
"""
家纺 希腊站上传表：按中国站分类树分类。
输入：input/希腊站上传_家纺.xlsx
输出：商品分类工作区/output/希腊站上传_家纺_已分类.xlsx
策略：很多行中文品名只有尺寸(如 60x180 / 270x280)，需结合希腊语品名判定类型
      (ΚΟΥΒΕΡΤΑ=毛毯 / ΠΑΤΑΚΙ ΔΑΠΕΔΟΥ=地垫 / ΣΕΤ ΣΕΝΤΟΝΙΩΝ=床单套 /
       ΑΔΙΑΒΡΟΧΟ ΚΑΛΥΜΜΑ=防水罩 / ΤΡΑΠΕΖΟΜΑΝΤΗΛΟ=桌布 / ΚΟΥΡΤΙΝΑ=窗帘)。
"""
import openpyxl, os
from collections import Counter

HERE = os.path.dirname(os.path.abspath(__file__))          # 商品分类工作区/
ROOT = os.path.dirname(HERE)                               # 仓库根
SRC = os.path.join(ROOT, "input", "希腊站上传_家纺.xlsx")
OUT = os.path.join(HERE, "output", "希腊站上传_家纺_已分类.xlsx")
TREE = os.path.join(HERE, "data", "中国站商品分类.xlsx")

twb = openpyxl.load_workbook(TREE, data_only=True); tws = twb.active
valid = set()
for r in range(1, tws.max_row + 1):
    a,b,c = tws.cell(r,4).value, tws.cell(r,5).value, tws.cell(r,6).value
    if a and b and c: valid.add((str(a),str(b),str(c)))

def classify(name, grk):
    g = '' if grk is None else str(grk).upper()
    # 窗帘（ΚΟΥΡΤΙΝΑ）优先，含纯尺寸行
    if 'ΚΟΥΡΤΙΝΑ' in g:
        return (None, ('家居百货','家纺系列','窗帘'))
    # 毛毯（ΚΟΥΒΕΡΤΑ / ΦΛΙΣ=羊羔绒毯）
    if 'ΚΟΥΒΕΡΤΑ' in g or 'ΦΛΙΣ' in g:
        return (None, ('家居百货','家纺系列','毛毯'))
    # 地垫（ΠΑΤΑΚΙ ΔΑΠΕΔΟΥ）
    if 'ΠΑΤΑΚΙ' in g or 'ΔΑΠΕΔΟΥ' in g or '地垫' in name:
        return (None, ('家居百货','家纺系列','地垫地毯'))
    # 床单套装（ΣΕΤ ΣΕΝΤΟΝΙΩΝ）：四件套/五件套
    if 'ΣΕΤ' in g or 'ΣΕΝΤΟΝΙ' in g or '床单' in name or '被套' in name:
        # 五件套(含被套)→ 近似床上四件套
        if '5' in name or '五件' in name or '5 ΤΕΜ' in g:
            return ('近似:床单五件套(含被套),中国站无专叶,归床上四件套', ('家居百货','家纺系列','床上四件套'))
        return (None, ('家居百货','家纺系列','床上四件套'))
    # 防水罩/床垫罩（ΑΔΙΑΒΡΟΧΟ ΚΑΛΥΜΜΑ ΣΤΡΩΜΑΤΟΣ，带+30松紧边=床笠）
    if 'ΑΔΙΑΒΡΟΧΟ' in g or 'ΚΑΛΥΜΜΑ' in g or '防水罩' in name or '床垫' in name:
        return (None, ('家居百货','家纺系列','床笠'))
    # 桌布（ΤΡΑΠΕΖΟΜΑΝΤΗΛΟ）
    if 'ΤΡΑΠΕΖΟΜΑΝΤΗΛΟ' in g or '桌布' in name:
        return (None, ('家居百货','家纺系列','桌布'))
    # 枕头套/枕套（ΜΑΞΙΛΑΡΟΘΗΚΗ=枕套 50x70）—— 需在"枕头"判定之前，避免误吞"枕头套"
    if 'ΜΑΞΙΛΑΡΟΘΗΚΗ' in g or '枕头套' in name or ('枕套' in name and '靠' not in name):
        return (None, ('家居百货','家纺系列','枕套'))
    # 靠枕（ΔΙΑΚΟΣΜΗΤΙΚΟ ΜΑΞΙΛΑΡΙ=装饰靠枕 45x45）≈ 抱枕
    if 'ΔΙΑΚΟΣΜΗΤΙΚΟ' in g or '靠枕' in name or '45x45' in name:
        return (None, ('家居百货','家纺系列','抱枕及抱枕套'))
    # 枕头（ΜΑΞΙΛΑΡΙ 50x70 睡眠枕成品）：中国站无"枕头"成品叶 → 近似抱枕及抱枕套
    if ('ΜΑΞΙΛΑΡΙ' in g or '枕头' in name) and 'ΘΗΚΗ' not in g:
        return ('近似:睡眠枕中国站无成品叶,归抱枕及抱枕套', ('家居百货','家纺系列','抱枕及抱枕套'))
    return None

wb = openpyxl.load_workbook(SRC); ws = wb.active
for mr in list(ws.merged_cells.ranges):
    if mr.min_col <= 5 and mr.max_col >= 4: ws.unmerge_cells(str(mr))

done = approx = uncls = 0; uncls_rows=[]; approx_rows=[]
for r in range(2, ws.max_row + 1):
    n = ws.cell(r,1).value
    if not n or not str(n).strip():
        ws.cell(r,5).value=ws.cell(r,6).value=ws.cell(r,7).value=None; continue
    name = str(n); g = ws.cell(r,2).value
    res = classify(name, g)
    if res is None:
        ws.cell(r,5).value=ws.cell(r,6).value=ws.cell(r,7).value=None
        uncls += 1; uncls_rows.append((r,name,g)); continue
    flag, chain = res
    if chain not in valid:
        ws.cell(r,5).value=ws.cell(r,6).value=ws.cell(r,7).value=None
        uncls += 1; uncls_rows.append((r,name,f"非法链{chain}")); continue
    ws.cell(r,5).value,ws.cell(r,6).value,ws.cell(r,7).value = chain
    if flag: approx += 1; approx_rows.append((r,name,flag,' / '.join(chain)))
    else: done += 1

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print(f"精确 {done} | 近似 {approx} | 待确认 {uncls}")
if approx_rows:
    print("近似:")
    for r,n,f,c in approx_rows: print(f"  {r:>3}| {n} ~ {c}  [{f}]")
if uncls_rows:
    print("待确认:")
    for row in uncls_rows: print("  ",row)
print("输出:", OUT)
