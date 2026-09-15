# -*- coding: utf-8 -*-
"""
fillback.py — 待审核商品分类「填回」主数据库（带确认闸门）
========================================================
背景:
  classify.py 把无法归类的商品标记为 系统复核=待确认（待审核）。
  用户在 Excel 里手动改好这些项后，本脚本把"曾被待审核、现已填好
  有效分类"的行写回 主数据库_商品数据.xlsx，下次自动命中（货号精确匹配）。

两阶段（由编排 bat 调用）:
  scan   分类后、用户编辑前：扫描 F2 中 系统复核=待确认 的行，存快照
  commit 用户关掉文件后：比对快照，找出已解决项 -> 列清单 -> 确认 -> 追加主数据库

依赖: 与 classify.py 同目录，复用其 norm / load_cn_tree / load_master / first_sheet。
"""
import sys
import os
import json
import argparse
import importlib.util

import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))
_spec = importlib.util.spec_from_file_location('cf', os.path.join(BASE, 'classify.py'))
cf = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(cf)

CN_TREE = cf.CN_TREE
MASTER = cf.MASTER


def norm(s):
    return cf.norm(s)


def find_class_cols(ws):
    """定位 F2 的最终分类列与系统复核列。"""
    header = [c.value for c in ws[1]]

    def idx(name):
        for i, h in enumerate(header):
            if h == name:
                return i
        return None

    t = idx('产品分类')
    s = idx('二级分类')
    l = idx('三级分类')
    rev = idx('系统复核')
    # 兜底：若精确列缺失（非标准结构），用 classify 的表头识别
    if t is None or s is None or l is None:
        cols = cf.locate_columns(ws)
        t, s, l = cols['otop'], cols['osec'], cols['oleaf']
    return t, s, l, rev


def compute_candidates(ws, pending, by_code, by_name, cn):
    """从已编辑的 F2 中找出『曾被待审核、现已填好有效分类』的候选行。
    pending: 编辑前快照里的待审核集合（{'code':..,'name':..}）。
    返回 (candidates, still_pending)：
      candidates   : [(name, code, bar, top, sec, leaf), ...] 已解决可填回
      still_pending: [(name, code, top, sec, leaf), ...]     仍待确认/无效链
    """
    pend_codes = {p['code'] for p in pending if p['code']}
    pend_names = {p['name'] for p in pending if p['name']}
    t, s, l, _ = find_class_cols(ws)
    candidates = []
    still_pending = []
    for row in ws.iter_rows(min_row=2):
        code = row[1].value if len(row) > 1 else None
        name = row[0].value if len(row) > 0 else None
        nk = norm(code)
        nn = norm(name)
        is_pending = (nk and nk in pend_codes) or (nn and nn in pend_names)
        if not is_pending:
            continue
        top = row[t].value if (t is not None and t < len(row)) else None
        sec = row[s].value if (s is not None and s < len(row)) else None
        leaf = row[l].value if (l is not None and l < len(row)) else None
        if leaf in (None, '待确认', '') or top in (None, '待确认', '') or sec in (None, '待确认', ''):
            still_pending.append((name, code, top, sec, leaf))
            continue
        chain = (top, sec, leaf)
        if chain not in cn:
            still_pending.append((name, code, top, sec, leaf))
            continue
        if (nk and nk in by_code) or (nn and nn in by_name):
            # 已在主数据库（可能之前填回过），不重复
            continue
        bar = row[2].value if len(row) > 2 else None
        candidates.append((name, code, bar, top, sec, leaf))
    return candidates, still_pending


def auto_apply_recommendations(classified, pending_codes, cn_chains):
    """把「程序推荐分类」sheet 里的推荐分类自动写回主表的「产品分类/二级/三级」列。
    触发条件（针对快照里的待审核行）：
      - 主表三列均空/待确认                → 自动采纳
      - 主表已填但整链不在中国站标准树里    → 自动采纳（极可能是用户输错/没填完）
      - 主表已填且整链有效                 → 不覆盖，尊重用户手动修改
    Returns [(name, code, top, sec, leaf), ...] 已采纳项；保存 F2。
    """
    wb = openpyxl.load_workbook(classified)
    ws = wb.worksheets[0]
    if '程序推荐分类' not in wb.sheetnames:
        wb.close()
        return []

    # 构建 货号norm -> (top, sec, leaf) 推荐映射（仅保留有效推荐）
    rec_map = {}
    ws_rec = wb['程序推荐分类']
    for row in ws_rec.iter_rows(min_row=2, values_only=True):
        if len(row) < 5:
            continue
        name, code, t, s, l = row[0], row[1], row[2], row[3], row[4]
        if code is None:
            continue
        nk = norm(code)
        if not nk:
            continue
        if t and s and l and t != '待确认' and s != '待确认' and l != '待确认':
            rec_map[nk] = (t, s, l)

    # 定位主表的「产品分类/二级分类/三级分类」列
    header = [c.value for c in ws[1]]

    def find_col(name):
        for i, h in enumerate(header):
            if h == name:
                return i
        return None
    t_idx = find_col('产品分类')
    s_idx = find_col('二级分类')
    l_idx = find_col('三级分类')
    if t_idx is None or s_idx is None or l_idx is None:
        wb.close()
        return []

    def is_empty_or_pending(v):
        return v is None or v == '待确认' or (isinstance(v, str) and v.strip() == '')

    applied = []
    for row in ws.iter_rows(min_row=2):
        if len(row) < 3:
            continue
        code = row[1].value
        nk = norm(code) if code else ''
        if nk not in pending_codes:
            continue
        if nk not in rec_map:
            continue
        t_new, s_new, l_new = rec_map[nk]
        t_cur = row[t_idx].value if t_idx < len(row) else None
        s_cur = row[s_idx].value if s_idx < len(row) else None
        l_cur = row[l_idx].value if l_idx < len(row) else None

        all_empty = (is_empty_or_pending(t_cur)
                     and is_empty_or_pending(s_cur)
                     and is_empty_or_pending(l_cur))
        cur_chain = (t_cur, s_cur, l_cur)
        cur_valid = cur_chain in cn_chains

        should_apply = all_empty or not cur_valid
        if not should_apply:
            continue  # 主表已是有效链，尊重用户手动修改

        row[t_idx].value = t_new
        row[s_idx].value = s_new
        row[l_idx].value = l_new
        applied.append((row[0].value, code, t_new, s_new, l_new))

    if applied:
        wb.save(classified)
        print(f"[auto-apply] 已自动采纳 {len(applied)} 条程序推荐分类到主表：")
        for (nm, cd, t, s, l) in applied:
            print(f"  - {nm} [{cd}] -> {t}/{s}/{l}")
    wb.close()
    return applied


def clean_for_upload(classified):
    """清除上传不需要的数据：删除辅助列「系统复核」（分类结果第34列，
    不在上传基准模板内，会触发表头严格校验失败）。原地覆盖保存。
    注：保留「程序推荐分类」辅助 sheet（上传脚本只读取首个 sheet，无影响）。"""
    wb = openpyxl.load_workbook(classified)
    ws = wb.worksheets[0]
    hdr = [c.value for c in ws[1]]
    if '系统复核' in hdr:
        col = hdr.index('系统复核') + 1
        ws.delete_cols(col, 1)
        print(f"[clean] 已删除辅助列「系统复核」(第{col}列)，文件回到标准 33 列模板")
    else:
        print("[clean] 未找到「系统复核」列，无需清理")
    wb.save(classified)


def scan(classified, snapshot_out):
    wb = openpyxl.load_workbook(classified, read_only=True, data_only=True)
    ws = cf.first_sheet(wb)
    _, _, _, rev = find_class_cols(ws)

    pending = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        is_pending = False
        if rev is not None and r[rev] == '待确认':
            is_pending = True
        # 兜底：无系统复核列时，按三级=待确认判定
        if not is_pending and rev is None and len(r) > 5 and r[5] == '待确认':
            is_pending = True
        if is_pending:
            code = r[1] if len(r) > 1 else None
            name = r[0] if len(r) > 0 else None
            pending.append({'code': norm(code) if code else '',
                            'name': norm(name) if name else ''})
    wb.close()

    with open(snapshot_out, 'w', encoding='utf-8') as f:
        json.dump(pending, f, ensure_ascii=False)
    print(f"[scan] 待审核扫描: {len(pending)} 行（系统复核=待确认）已记录到快照")
    print(f"       快照: {snapshot_out}")


def commit(classified, snapshot_in, master=None, tree=None):
    master = master or MASTER
    tree = tree or CN_TREE
    cn = cf.load_cn_tree(tree)
    by_code, by_name = cf.load_master(master)
    print(f"[commit] 主数据库现有: 货号{len(by_code)} / 品名{len(by_name)}")

    wb = openpyxl.load_workbook(classified, data_only=True)
    ws = cf.first_sheet(wb)
    try:
        pending = json.load(open(snapshot_in, encoding='utf-8'))
    except Exception as e:
        print(f"[错误] 读取快照失败: {e}")
        sys.exit(1)
    candidates, still_pending = compute_candidates(ws, pending, by_code, by_name, cn)
    wb.close()

    for (nm, cd, tp, se, lf) in still_pending:
        print(f"  [跳过-未解决] {nm} ({cd}) 仍为: {tp}/{se}/{lf}")
    if not candidates:
        print("[commit] 没有可填回的新分类（待审核项未解决 / 均已在库）。")
    else:
        print(f"\n[commit] 发现 {len(candidates)} 条待审核分类已解决，拟填回主数据库：")
        for (nm, cd, _b, tp, se, lf) in candidates:
            print(f"  - {nm} [{cd}] -> {tp}/{se}/{lf}")
        ans = input("确认填回主数据库？(Y/N): ").strip().lower()
        if ans not in ('y', 'yes', '是'):
            print("[commit] 已取消填回，主数据库保持不变。")
            return
        mb = openpyxl.load_workbook(master)
        mws = mb.worksheets[0]
        for (nm, cd, bar, tp, se, lf) in candidates:
            mws.append([nm, cd, bar, tp, se, lf])
        mb.save(master)
        print(f"[commit] 已填回 {len(candidates)} 条到主数据库: {master}")
    # 无论是否填回，都清理上传辅助列（待确认按钮流程的兜底路径）
    clean_for_upload(classified)


def prepare(classified, snapshot_in, out_json, master=None, tree=None):
    """计算待审核填回候选，输出 JSON 供 GUI 确认窗口展示。
    步骤：1) 先从「程序推荐分类」sheet 自动采纳到主表（针对待审核行），
          2) 再按最新主表计算「已解决/仍待确认」候选，
          3) 写 JSON 含三类信息：resolved / unresolved / auto_applied。
    不写主数据库、不清理文件（确认后才由 apply 执行）。"""
    master = master or MASTER
    tree = tree or CN_TREE
    cn = cf.load_cn_tree(tree)
    by_code, by_name = cf.load_master(master)

    try:
        pending = json.load(open(snapshot_in, encoding='utf-8'))
    except Exception as e:
        print(f"[错误] 读取快照失败: {e}")
        sys.exit(1)

    pend_codes = {p['code'] for p in pending if p['code']}

    # 1) 自动采纳推荐分类到主表（修改 F2 文件）
    auto_applied = auto_apply_recommendations(classified, pend_codes, cn)

    # 2) 重新读取 F2，按最新主表计算候选
    wb = openpyxl.load_workbook(classified, data_only=True)
    ws = cf.first_sheet(wb)
    candidates, still_pending = compute_candidates(ws, pending, by_code, by_name, cn)
    wb.close()

    resolved = [{'name': nm, 'code': cd, 'chain': f"{tp}/{se}/{lf}"}
                for (nm, cd, _b, tp, se, lf) in candidates]
    unresolved = [{'name': nm, 'code': cd, 'current': f"{tp}/{se}/{lf}"}
                  for (nm, cd, tp, se, lf) in still_pending]
    auto_json = [{'name': nm, 'code': cd, 'chain': f"{t}/{s}/{l}"}
                 for (nm, cd, t, s, l) in auto_applied]
    payload = {'resolved': resolved, 'unresolved': unresolved, 'auto_applied': auto_json}
    with open(out_json, 'w', encoding='utf-8') as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    print(f"[prepare] 自动采纳 {len(auto_applied)} 条 | 已解决 {len(resolved)} 条 | 仍待确认 {len(unresolved)} 条")
    print(f"          候选清单已写入: {out_json}")


def apply(classified, snapshot_in, master=None, tree=None):
    """确认后执行：先自动采纳推荐分类 → 把『已解决』待审核项填回主数据库 → 清理上传辅助列。
    由一键流程在 GUI 确认(yes)后调用。"""
    master = master or MASTER
    tree = tree or CN_TREE
    cn = cf.load_cn_tree(tree)
    by_code, by_name = cf.load_master(master)

    try:
        pending = json.load(open(snapshot_in, encoding='utf-8'))
    except Exception as e:
        print(f"[错误] 读取快照失败: {e}")
        sys.exit(1)

    pend_codes = {p['code'] for p in pending if p['code']}

    # 1) 自动采纳（GUI 已确认；此处再跑一次保证幂等）
    auto_applied = auto_apply_recommendations(classified, pend_codes, cn)

    # 2) 重新读取最新主表
    wb = openpyxl.load_workbook(classified, data_only=True)
    ws = cf.first_sheet(wb)
    candidates, still_pending = compute_candidates(ws, pending, by_code, by_name, cn)
    wb.close()

    for (nm, cd, tp, se, lf) in still_pending:
        print(f"  [提示-未解决] {nm} ({cd}) 仍为: {tp}/{se}/{lf}（保留原样，可能上传报错）")
    if candidates:
        print(f"[apply] 填回 {len(candidates)} 条到主数据库：")
        for (nm, cd, _b, tp, se, lf) in candidates:
            print(f"  - {nm} [{cd}] -> {tp}/{se}/{lf}")
        mb = openpyxl.load_workbook(master)
        mws = mb.worksheets[0]
        for (nm, cd, bar, tp, se, lf) in candidates:
            mws.append([nm, cd, bar, tp, se, lf])
        mb.save(master)
        print(f"[apply] 已填回主数据库: {master}")
    else:
        print("[apply] 没有可填回的新分类（待审核项未解决 / 均已在库）。")
    clean_for_upload(classified)
    print("[apply] 上传辅助列已清理，可以继续上传。")


def main():
    ap = argparse.ArgumentParser(description="待审核商品分类填回主数据库")
    sub = ap.add_subparsers(dest='cmd', required=True)

    sp = sub.add_parser('scan', help='扫描待审核行并存快照')
    sp.add_argument('--classified', required=True)
    sp.add_argument('--out', required=True)

    cp = sub.add_parser('commit', help='比对快照，确认(Y/N)后填回主数据库')
    cp.add_argument('--classified', required=True)
    cp.add_argument('--snapshot', required=True)
    cp.add_argument('--master', default=MASTER)
    cp.add_argument('--tree', default=CN_TREE)

    pp = sub.add_parser('prepare', help='计算待审核填回候选，输出 JSON 供 GUI 确认')
    pp.add_argument('--classified', required=True)
    pp.add_argument('--snapshot', required=True)
    pp.add_argument('--out', required=True)
    pp.add_argument('--master', default=MASTER)
    pp.add_argument('--tree', default=CN_TREE)

    ap2 = sub.add_parser('apply', help='确认后：填回主数据库 + 清理上传辅助列')
    ap2.add_argument('--classified', required=True)
    ap2.add_argument('--snapshot', required=True)
    ap2.add_argument('--master', default=MASTER)
    ap2.add_argument('--tree', default=CN_TREE)

    args = ap.parse_args()
    if args.cmd == 'scan':
        scan(args.classified, args.out)
    elif args.cmd == 'commit':
        commit(args.classified, args.snapshot, args.master, args.tree)
    elif args.cmd == 'prepare':
        prepare(args.classified, args.snapshot, args.out, args.master, args.tree)
    elif args.cmd == 'apply':
        apply(args.classified, args.snapshot, args.master, args.tree)


if __name__ == '__main__':
    main()
