# -*- coding: utf-8 -*-
import openpyxl, collections, os
BASE = os.path.dirname(os.path.abspath(__file__))
p = r'<原始分类表路径>\中国站商品分类 - 副本.xlsx'   # 换成你自己的路径
wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
ws = wb['Sheet1']
topid = collections.OrderedDict()  # 一级ID前3位 -> 一级名
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[0] is None:
        continue
    tid = str(r[0])[:3]
    tname = r[3]
    if tid not in topid:
        topid[tid] = tname
ws2 = wb['Sheet2']
enabled = [r[0] for r in ws2.iter_rows(values_only=True) if r[0]]
en = set(enabled)
print('一级ID段 -> 一级名 -> 启用?')
disabled_ids = []
for tid in sorted(topid, key=lambda x: int(x)):
    nm = topid[tid]
    ok = nm in en
    if not ok:
        disabled_ids.append(tid)
    print(f'  {tid}: {nm}  {"启用" if ok else "禁用"}')
print()
print('禁用一级ID段:', disabled_ids)
wb.close()
