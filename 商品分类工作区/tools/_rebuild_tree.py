# -*- coding: utf-8 -*-
"""用《中国站商品分类 - 副本.xlsx》的 Sheet1 重建权威树 data/中国站商品分类.xlsx，
规则：以"一级分类 ID 前3位"判定启用/禁用（白名单在 Sheet2）。
禁用一级 ID 段（不在白名单的一级）下的所有链（无论二级/三级多精准）一律剔除。
输出：data/中国站商品分类.xlsx  (列: 一级ID,二级ID,三级ID,产品分类,二级分类,三级分类)
"""
import openpyxl, os, collections

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = r'<原始分类表路径>\中国站商品分类 - 副本.xlsx'   # 换成你自己的路径
OUT = os.path.join(ROOT, 'data', '中国站商品分类.xlsx')

wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

# 1) 一级ID前3位 -> 一级名（来自 Sheet1）
topid = collections.OrderedDict()
for r in wb['Sheet1'].iter_rows(min_row=2, values_only=True):
    if r[0] is None:
        continue
    tid = str(r[0])[:3]
    if tid not in topid:
        topid[tid] = r[3]

# 2) 白名单（Sheet2，只有一级名）
enabled = set(r[0] for r in wb['Sheet2'].iter_rows(values_only=True) if r[0])

# 3) 禁用一级 ID 段
disabled_seg = [tid for tid in topid if topid[tid] not in enabled]
print('一级名->ID段:', dict(topid))
print('禁用一级ID段:', disabled_seg, '=>', [topid[s] for s in disabled_seg])

# 4) 重建：只保留一级ID前3位不在禁用段 的链
out = openpyxl.Workbook()
ws = out.active
ws.title = 'Sheet1'
ws.append(['一级ID', '二级ID', '三级ID', '产品分类', '二级分类', '三级分类'])

kept = dropped = 0
for r in wb['Sheet1'].iter_rows(min_row=2, values_only=True):
    if r[0] is None or r[3] is None or r[4] is None or r[5] is None:
        continue
    seg = str(r[0])[:3]
    if seg in disabled_seg:
        dropped += 1
        continue
    ws.append([r[0], r[1], r[2], r[3], r[4], r[5]])
    kept += 1
wb.close()
out.save(OUT)
print(f'写出 {OUT}: 保留链={kept} 剔除(禁用段)={dropped}')

# 5) 自检：重建后的树不应含任何禁用一级
wb2 = openpyxl.load_workbook(OUT, read_only=True, data_only=True)
chains = set()
bad = []
for r in wb2['Sheet1'].iter_rows(min_row=2, values_only=True):
    if r[3] and r[4] and r[5]:
        chains.add((r[3], r[4], r[5]))
        if topid.get(str(r[0])[:3]) in (topid[s] for s in disabled_seg):
            bad.append(r[3])
print(f'新树链数={len(chains)} 含禁用一级={len(bad)}')
wb2.close()
