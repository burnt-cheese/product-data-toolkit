# -*- coding: utf-8 -*-
"""
厨房小家电(厨房/个护小家电) 希腊站上传表：按中国站分类树分类。
输入：input/希腊站上传_厨房小家电.xlsx
输出：output/希腊站上传_厨房小家电_已分类.xlsx
数据：真实商品行 2..47（行 48+ 为空），共 46 个。
"""
import openpyxl, os, re
from collections import Counter

HERE = os.path.dirname(os.path.abspath(__file__))          # 商品分类工作区/
ROOT = os.path.dirname(HERE)                               # 仓库根
SRC = os.path.join(ROOT, "input", "希腊站上传_厨房小家电.xlsx")
TREE = os.path.join(HERE, "data", "中国站商品分类.xlsx")
OUT = os.path.join(HERE, "output", "希腊站上传_厨房小家电_已分类.xlsx")

# ---------- 1. 中国站分类树合法 3 级链路 ----------
twb = openpyxl.load_workbook(TREE, data_only=True)
tws = twb.active
valid = set()
for r in range(1, tws.max_row + 1):
    a, b, c = tws.cell(r, 4).value, tws.cell(r, 5).value, tws.cell(r, 6).value
    if a and b and c:
        valid.add((str(a), str(b), str(c)))

# ---------- 2. 分类规则（基于“品名”关键词，顺序：具体→泛） ----------
def classify(n):
    # ---- 电动煎烤平板（三明治机/接触式烤盘）----
    if '三明治机' in n:
        return ('家电数码', '厨房电器', '煎烤器')
    if '接触式烤盘' in n or '烤盘' in n:
        return ('家电数码', '厨房电器', '煎烤器')
    # ---- 吐司机（烤面包机）----
    if '烤面包机' in n:
        return ('家电数码', '厨房电器', '吐司机')
    # ---- 咖啡 ----
    if '咖啡机' in n:
        return ('家电数码', '厨房电器', '咖啡机')
    if '咖啡壶' in n:
        return ('家电数码', '厨房电器', '咖啡壶')
    # ---- 水壶 ----
    if '电热水壶' in n:
        return ('家电数码', '厨房电器', '电热水壶')
    if '热水壶' in n or '电水壶' in n:
        return ('家电数码', '厨房电器', '电热水壶')
    # ---- 烤箱 ----
    if '烤箱' in n:
        return ('家电数码', '厨房电器', '烤箱')
    # ---- 电磁炉 ----
    if '电磁炉' in n:
        return ('家电数码', '厨房电器', '电磁炉')
    # ---- 绞肉机 ----
    if '绞肉机' in n:
        return ('家电数码', '厨房电器', '绞肉机')
    # ---- 真空（机器 vs 袋耗材）----
    # 真空封口机预制袋/袋卷 → 真空袋耗材（非机器）
    if '真空' in n and ('袋' in n):
        return ('家居百货', '家居收纳', '真空袋')
    # ---- 打蛋器（电动）----
    if '打蛋器' in n:
        return ('家电数码', '厨房电器', '电动打蛋器')
    # ---- 搅拌机（含台式/破壁/手持料理棒/冰沙机，均归搅拌机）----
    if '搅拌机' in n or '搅拌器' in n or '料理棒' in n or '破壁' in n:
        return ('家电数码', '厨房电器', '搅拌机')
    # ---- 榨汁机（含慢磨/柑橘）----
    if '榨汁机' in n or '慢磨' in n or '榨汁' in n:
        return ('家电数码', '厨房电器', '榨汁机')
    # ---- 煮蛋器 ----
    if '煮蛋器' in n:
        return ('家电数码', '厨房电器', '蒸蛋器')
    # ---- 电子秤（厨房秤/浴秤/体脂秤）----
    if '电子秤' in n or '浴室秤' in n or '体脂秤' in n or '厨房电子秤' in n:
        return ('家电数码', '电子产品', '电子秤')
    # ---- 熨斗 ----
    if '熨斗' in n:
        return ('家电数码', '小家电', '电熨斗')
    # ---- 美发 ----
    if '直发器' in n:
        return ('家电数码', '个人护理电器', '夹板')
    if '卷发器' in n:
        return ('家电数码', '个人护理电器', '卷发棒')
    # ---- 理容 ----
    if '电推剪' in n:
        return ('家电数码', '个人护理电器', '理发剪')
    if '修剪器' in n:
        return ('家电数码', '个人护理电器', '剃毛器')
    return None

# ---------- 3. 处理主表（只处理真实商品行，行 48+ 为空跳过） ----------
wb = openpyxl.load_workbook(SRC)
ws = wb.active
# 模板 D:E（条形码+产品分类）按行合并，先解除再写 产品分类(E)/二级(F)/三级(G)
for mr in list(ws.merged_cells.ranges):
    if mr.min_col <= 5 and mr.max_col >= 4:
        ws.unmerge_cells(str(mr))

class_log, invalid, empty_rows = [], [], []
for r in range(2, ws.max_row + 1):
    name = ws.cell(r, 1).value
    if not name or not str(name).strip():
        # 空行：清空分类列（避免模板残留）并跳过
        ws.cell(r, 5).value = ws.cell(r, 6).value = ws.cell(r, 7).value = None
        empty_rows.append(r)
        continue
    name = str(name)
    chain = classify(name)
    if chain is None:
        class_log.append((r, name, '待确认'))
        continue
    if chain not in valid:
        invalid.append((r, name, chain))
        class_log.append((r, name, '链路不存在(待确认)'))
        continue
    ws.cell(r, 5).value, ws.cell(r, 6).value, ws.cell(r, 7).value = chain
    class_log.append((r, name, ' / '.join(chain)))

# ---------- 4. 落盘 & 报告 ----------
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)

print("=== 非法链路 (应为 0) ===", invalid if invalid else "无")
print("\n=== 分类分布 ===")
cnt = Counter(c for r, n, c in class_log)
for c, k in sorted(cnt.items()):
    print(f"  {k:>3}  {c}")
print("\n=== 未分类明细 ===")
for r, n, c in class_log:
    if '待确认' in c:
        print(f"  行{r:>3}: {n!r}")
print("\n=== 汇总 ===")
filled = sum(1 for r, n, c in class_log if '待确认' not in c)
print("真实商品行数:", len(class_log))
print("已分类:", filled)
print("待确认:", len(class_log) - filled)
print("空行(跳过):", len(empty_rows))
print("输出:", OUT)
