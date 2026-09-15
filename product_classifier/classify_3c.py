# -*- coding: utf-8 -*-
"""
3C数码 希腊站上传表：按中国站分类树分类（3C 数码配件）。
输入：input/希腊站上传_3C数码.xlsx
输出：output/希腊站上传_3C数码_classified.xlsx
"""
import openpyxl, os
from collections import Counter

HERE = os.path.dirname(os.path.abspath(__file__))          # product_classifier/
ROOT = os.path.dirname(HERE)                               # 仓库根
SRC = os.path.join(ROOT, "input", "希腊站上传_3C数码.xlsx")
TREE = os.path.join(HERE, "data", "中国站商品分类.xlsx")
OUT = os.path.join(HERE, "output", "希腊站上传_3C数码_classified.xlsx")

# ---------- 1. 中国站分类树合法 3 级链路 ----------
twb = openpyxl.load_workbook(TREE, data_only=True)
tws = twb.active
valid = set()
for r in range(1, tws.max_row + 1):
    a, b, c = tws.cell(r, 4).value, tws.cell(r, 5).value, tws.cell(r, 6).value
    if a and b and c:
        valid.add((str(a), str(b), str(c)))

# ---------- 2. 分类规则（基于“品名”关键词） ----------
def classify(n):
    # 音箱
    if '蓝牙音箱' in n:
        return ('家电数码', '数码影音', '便携音响')
    # 麦克风 / 话筒
    if '麦克风' in n:
        return ('家电数码', '直播工具', '话筒')
    # 键鼠套装
    if '键鼠套装' in n:
        return ('家电数码', '电脑配件', '键鼠套装')
    # 鼠标垫
    if '鼠标垫' in n:
        return ('家电数码', '电脑配件', '鼠标垫')
    # 鼠标
    if '鼠标' in n:
        return ('家电数码', '电脑配件', '鼠标')
    # 键盘
    if '键盘' in n:
        return ('家电数码', '电脑配件', '键盘')
    # 音频线
    if '音频线' in n:
        return ('家电数码', '手机配件', '音频线')
    # 数据线
    if '数据线' in n:
        return ('家电数码', '手机配件', '数据线')
    # 线缆（无详情，近似归 延长线，需人工确认）
    if '线缆' in n:
        return ('家电数码', '家用电器', '延长线')
    # 车载 FM 发射器（无专属叶，近似 车载蓝牙，需人工确认）
    if 'FM发射器' in n or 'FM 发射器' in n:
        return ('家电数码', '电子产品', '车载蓝牙')
    # 车载充电器
    if '车载充电器' in n:
        return ('家电数码', '手机件', '车载充电器')
    # 墙充
    if '充电器' in n:
        return ('家电数码', '手机配件', '充电器')
    # 摩托手机支架
    if '摩托车手机支架' in n:
        return ('骑行用品', '摩托车功能件', '手机支架')
    # 车载/后排手机支架
    if '车载手机支架' in n or '车载后排手机支架' in n:
        return ('家电数码', '手机配件', '汽车手机支架')
    # 普通手机支架
    if '手机支架' in n:
        return ('家电数码', '手机配件', '手机支架')
    # 蓝牙耳机（含 无线耳机 / 真无线 / TWS / OWS / ANC / ENC / AI）
    # 注意“真无线降噪耳机”中“无线”与“耳机”不连续，故用 '无线' 而非 '无线耳机'；
    # 其它含“无线”的产品（无线鼠标/键鼠套装/麦克风）已在前面规则先拦截，不会误入此处。
    if '蓝牙耳机' in n or '无线' in n or 'TWS' in n or 'OWS' in n \
       or 'ANC' in n or 'ENC' in n:
        return ('家电数码', '手机配件', '蓝牙耳机')
    # 有线耳机
    if '有线耳机' in n:
        return ('家电数码', '手机配件', '耳机')
    # 兜底耳机
    if '耳机' in n:
        return ('家电数码', '手机配件', '耳机')
    return None

# ---------- 3. 处理主表 ----------
wb = openpyxl.load_workbook(SRC)
ws = wb.active
# 模板 D:E（条形码+产品分类）按行合并，先解除再写 产品分类(E)/二级(F)/三级(G)
for mr in list(ws.merged_cells.ranges):
    if mr.min_col <= 5 and mr.max_col >= 4:
        ws.unmerge_cells(str(mr))

class_log, invalid, need_review = [], [], []
for r in range(2, ws.max_row + 1):
    name = ws.cell(r, 1).value or ''
    chain = classify(name)
    if chain is None:
        class_log.append((r, name, '待确认'))
        continue
    if chain not in valid:
        invalid.append((r, name, chain))
        class_log.append((r, name, '链路不存在(待确认)'))
        continue
    ws.cell(r, 5).value, ws.cell(r, 6).value, ws.cell(r, 7).value = chain
    class_log.append((r, name, ' / '.join(chain)))
    if chain[2] in ('延长线', '车载蓝牙'):
        need_review.append((r, name, ' / '.join(chain)))

# ---------- 4. 落盘 & 报告 ----------
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)

print("=== 非法链路 (应为 0) ===", invalid if invalid else "无")
print("\n=== 分类分布 ===")
cnt = Counter(c for r, n, c in class_log)
for c, k in sorted(cnt.items()):
    print(f"  {k:>3}  {c}")
print("\n=== 需人工确认（近似映射）===")
for r, n, c in need_review:
    print(f"  行{r:>3}: {n!r} -> {c}")
print("\n=== 汇总 ===")
print("总行数:", ws.max_row - 1)
print("已分类:", sum(1 for r, n, c in class_log if '待确认' not in c))
print("待确认:", sum(1 for r, n, c in class_log if '待确认' in c))
print("输出:", OUT)
