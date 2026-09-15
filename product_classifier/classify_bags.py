# -*- coding: utf-8 -*-
"""
箱包 希腊站上传表：按中国站分类树分类（80 个手提女包）。
用户口径：这批包(TΣΑΝΤΑ)仅泛称"包"+型号+颜色，无包型信息，统一归 箱包系列/女包/手提包。
输入：input/希腊站上传_箱包.xlsx
输出：product_classifier/output/希腊站上传_箱包_classified.xlsx
"""
import openpyxl, os
from collections import Counter

HERE = os.path.dirname(os.path.abspath(__file__))          # product_classifier/
ROOT = os.path.dirname(HERE)                               # 仓库根
SRC = os.path.join(ROOT, "input", "希腊站上传_箱包.xlsx")
OUT = os.path.join(HERE, "output", "希腊站上传_箱包_classified.xlsx")
TREE = os.path.join(HERE, "data", "中国站商品分类.xlsx")

twb = openpyxl.load_workbook(TREE, data_only=True); tws = twb.active
valid = set()
for r in range(1, tws.max_row + 1):
    a,b,c = tws.cell(r,4).value, tws.cell(r,5).value, tws.cell(r,6).value
    if a and b and c: valid.add((str(a),str(b),str(c)))

CHAIN = ('箱包系列', '女包', '手提包')
assert CHAIN in valid, "链路不存在!"

wb = openpyxl.load_workbook(SRC); ws = wb.active
for mr in list(ws.merged_cells.ranges):
    if mr.min_col <= 5 and mr.max_col >= 4: ws.unmerge_cells(str(mr))

done = skip = 0
for r in range(2, ws.max_row + 1):
    n = ws.cell(r,1).value
    if not n or not str(n).strip():
        ws.cell(r,5).value=ws.cell(r,6).value=ws.cell(r,7).value=None
        continue
    ws.cell(r,5).value, ws.cell(r,6).value, ws.cell(r,7).value = CHAIN
    done += 1

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print(f"已分类 {done} 行 -> {' / '.join(CHAIN)}")
# verify
wb=openpyxl.load_workbook(OUT); ws=wb.active
total=sum(1 for r in range(2,ws.max_row+1) if ws.cell(r,1).value)
full=sum(1 for r in range(2,ws.max_row+1) if ws.cell(r,7).value)
print(f"非空品名 {total} | 全三级 {full} | 残留D:E合并 {[str(m) for m in ws.merged_cells.ranges if m.min_col<=5 and m.max_col>=4]}")
print("输出:", OUT)
